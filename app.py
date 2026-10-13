import os
import io
import json
import textwrap
import time
from datetime import datetime

import streamlit as st
import pandas as pd
from dotenv import load_dotenv
from openai import OpenAI
from pypdf import PdfReader
from docx import Document


load_dotenv()

AZURE_OPENAI_ENDPOINT = os.getenv("AZURE_OPENAI_ENDPOINT", "").rstrip("/")
AZURE_OPENAI_API_KEY = os.getenv("AZURE_OPENAI_API_KEY", "")
AZURE_OPENAI_MODEL = os.getenv("AZURE_OPENAI_MODEL", "gpt-5.4-mini")

SUPPORTED_FILE_TYPES = ["pdf", "docx", "txt"]
MAX_CHARS_PER_DOCUMENT = 45000


st.set_page_config(
    page_title="Forvis Mazars Contracts Analyser",
    page_icon="📄",
    layout="wide"
)



if "uploader_key" not in st.session_state:
    st.session_state.uploader_key = 0




st.markdown("""
<style>

/* Main Title */
.main-title {
    font-size: 42px;
    font-weight: 800;
    color: #4169E1;
    margin-bottom: 5px;
    animation: fadeInDown 0.8s ease-out;
}

/* Subtitle */
.subtitle {
    font-size: 17px;
    color: #94A3B8;
    margin-bottom: 25px;
    animation: fadeInUp 1s ease-out;
}

/* Information Box */
.info-box {
    background: linear-gradient(
        135deg,
        #1E293B 0%,
        #334155 100%
    );
    border: 1px solid rgba(255,255,255,0.1);
    padding: 18px;
    border-radius: 16px;
    margin-bottom: 20px;
    color: white;
}

/* Metric Cards */
.metric-card {
    background: white;
    border: 1px solid #E2E8F0;
    padding: 16px;
    border-radius: 16px;
    transition: all 0.3s ease;
    box-shadow: 0 4px 12px rgba(0,0,0,0.05);
}

/* Hover Effect */
.metric-card:hover {
    transform: translateY(-5px);
    box-shadow: 0 15px 30px rgba(65,105,225,0.15);
    border-color: #4169E1;
}
            


/* Link Button */
[data-testid="stLinkButton"] a {
    transition: all 0.3s ease !important;
    border-radius: 12px !important;
}

/* Hover Animation */
[data-testid="stLinkButton"] a:hover {
    transform: scale(1.04);
    box-shadow: 0 10px 25px rg*a(65, 105, 225, 0.35);
}



/* Sidebar */
[data-testid="stSidebar"] {
    background: linear-gradient(
        180deg,
        #818589 0%,
        #003A70 100%
    );
}

[data-testid="stSidebar"] * {
    color: white !important;
}

/* Buttons */
.stButton button {
    width: 100%;
    background: linear-gradient(
        90deg,
        #4169E1 0%,
        #6A5ACD 100%
    );
    color: white;
    border: none;
    border-radius: 12px;
    padding: 10px 20px;
    font-weight: 600;
    transition: all 0.3s ease;
}

.stButton button:hover {
    transform: scale(1.04);
}

/* File Uploader */
[data-testid="stFileUploader"] {
    border-radius: 16px;
    border: 2px dashed #4169E1;
    padding: 20px;
}

/* Tabs */
.stTabs [role="tab"] {
    border-radius: 12px;
    padding: 10px 20px;
}

.stTabs [aria-selected="true"] {
    background-color: #4169E1;
    color: white;
}

/* Animations */

@keyframes fadeInDown {
    from {
        opacity: 0;
        transform: translateY(-20px);
    }
    to {
        opacity: 1;
        transform: translateY(0);
    }
}

@keyframes fadeInUp {
    from {
        opacity: 0;
        transform: translateY(20px);
    }
    to {
        opacity: 1;
        transform: translateY(0);
    }
}

</style>
""", unsafe_allow_html=True)


st.logo("Files/logoMazars.png")

st.markdown("<div class='main-title'>Forvis Mazars contract analyser</div>", unsafe_allow_html=True)
st.markdown(
    "<div class='subtitle'>Upload contracts, compare documents, detect missing clauses, review risks, chat with documents, and download Word or Excel outputs.</div>",
    unsafe_allow_html=True
)


def check_environment():
    missing = []

    if not AZURE_OPENAI_ENDPOINT:
        missing.append("AZURE_OPENAI_ENDPOINT")

    if not AZURE_OPENAI_API_KEY:
        missing.append("AZURE_OPENAI_API_KEY")

    if not AZURE_OPENAI_MODEL:
        missing.append("AZURE_OPENAI_MODEL")

    return missing


@st.cache_resource
def get_openai_client():
    return OpenAI(
        api_key=AZURE_OPENAI_API_KEY,
        base_url=AZURE_OPENAI_ENDPOINT
    )


def extract_text_from_pdf(uploaded_file):
    reader = PdfReader(uploaded_file)
    text_parts = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
            text_parts.append(f"\n--- Page {page_number} ---\n{page_text}")
        except Exception as error:
            text_parts.append(f"\n--- Page {page_number} ---\nCould not extract text: {error}")

    return "\n".join(text_parts).strip()


def extract_text_from_docx(uploaded_file):
    document = Document(uploaded_file)
    text_parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text.strip())

    for table_number, table in enumerate(document.tables, start=1):
        text_parts.append(f"\n--- Table {table_number} ---")
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                text_parts.append(row_text)

    return "\n".join(text_parts).strip()


def extract_text_from_txt(uploaded_file):
    raw_data = uploaded_file.read()

    try:
        return raw_data.decode("utf-8")
    except UnicodeDecodeError:
        return raw_data.decode("latin-1", errors="ignore")


def extract_text(uploaded_file):
    file_name = uploaded_file.name
    file_extension = file_name.split(".")[-1].lower()

    if file_extension == "pdf":
        text = extract_text_from_pdf(uploaded_file)
    elif file_extension == "docx":
        text = extract_text_from_docx(uploaded_file)
    elif file_extension == "txt":
        text = extract_text_from_txt(uploaded_file)
    else:
        raise ValueError(f"Unsupported file type: {file_extension}")

    if not text.strip():
        text = "No readable text could be extracted from this document."

    if len(text) > MAX_CHARS_PER_DOCUMENT:
        text = text[:MAX_CHARS_PER_DOCUMENT]
        text += "\n\nDocument was truncated because it exceeded the configured character limit."

    return {
        "file_name": file_name,
        "text": text
    }


def build_documents_context(documents):
    document_sections = []

    for index, document in enumerate(documents, start=1):
        section = f"""
Document {index}
File name: {document["file_name"]}

Content:
{document["text"]}
"""
        document_sections.append(textwrap.dedent(section))

    return "\n\n".join(document_sections)


def call_model(prompt, instructions):
    client = get_openai_client()

    response = client.responses.create(
        model=AZURE_OPENAI_MODEL,
        instructions=instructions,
        input=prompt
    )

    if hasattr(response, "output_text") and response.output_text:
        return response.output_text

    try:
        output_text = []
        for output_item in response.output:
            if hasattr(output_item, "content"):
                for content_item in output_item.content:
                    if hasattr(content_item, "text"):
                        output_text.append(content_item.text)

        return "\n".join(output_text).strip()
    except Exception:
        return str(response)


def safely_parse_json(raw_text):
    try:
        return json.loads(raw_text)
    except Exception:
        pass

    try:
        start_index = raw_text.find("{")
        end_index = raw_text.rfind("}") + 1
        json_text = raw_text[start_index:end_index]
        return json.loads(json_text)
    except Exception:
        return None


def compare_documents(documents):
    documents_context = build_documents_context(documents)

    prompt = f"""
You are a senior contracts review assistant for a professional services firm.

Compare the uploaded contract documents carefully and produce a structured comparison.

Your output must include:

1. Executive Summary
2. Document Overview
3. Key Differences Table
4. Missing or Added Clauses
5. Risk Assessment
6. Suggested Negotiation Points
7. Recommended Next Steps

Important instructions:
Do not invent clauses that are not present in the documents.
If extracted text is incomplete, clearly say so.
Use plain professional English.
This is an AI-assisted comparison and must be reviewed by a qualified human reviewer.

Uploaded documents:

{documents_context}
"""

    return call_model(
        prompt=prompt,
        instructions="You are a careful, precise and professional contracts comparison assistant."
    )


def analyse_missing_clauses_and_risks(documents):
    documents_context = build_documents_context(documents)

    prompt = f"""
You are a contracts risk and clause review assistant.

Analyse the uploaded documents and return ONLY valid JSON.

The JSON must use this exact structure:

{{
  "summary": {{
    "overall_risk_level": "High or Medium or Low",
    "overall_comments": "Short summary"
  }},
  "missing_clauses": [
    {{
      "clause_name": "Clause name",
      "status": "Missing or Present or Partial",
      "document_reference": "Document name or all documents",
      "impact": "Why this matters",
      "recommended_action": "What should be done"
    }}
  ],
  "risk_register": [
    {{
      "risk_area": "Risk area",
      "clause_reference": "Clause or section reference if available",
      "risk_level": "High or Medium or Low",
      "issue_identified": "Issue identified",
      "business_impact": "Business or legal impact",
      "recommended_action": "Recommended action"
    }}
  ]
}}

Review these common contract areas:
Confidentiality
Data protection
Liability cap
Indemnity
Termination
Payment terms
Intellectual property
Governing law
Dispute resolution
Audit rights
Service levels
Renewal
Notice period
Subcontracting
Assignment
Non-solicitation
Force majeure
Compliance with laws
Information security
Data breach notification

Important:
Return only JSON.
Do not include markdown.
Do not include explanations outside JSON.
Do not invent clauses that are not present.
If a clause cannot be confirmed, mark it as Partial or Missing.

Uploaded documents:

{documents_context}
"""

    raw_output = call_model(
        prompt=prompt,
        instructions="You are a contracts risk analyst. Return only valid JSON."
    )

    parsed_output = safely_parse_json(raw_output)

    if parsed_output is None:
        parsed_output = {
            "summary": {
                "overall_risk_level": "Unknown",
                "overall_comments": "The AI response could not be parsed into structured JSON."
            },
            "missing_clauses": [],
            "risk_register": []
        }

    return parsed_output, raw_output


def chat_with_documents(question, documents, chat_history):
    documents_context = build_documents_context(documents)

    history_text = ""
    for message in chat_history[-6:]:
        role = message.get("role", "")
        content = message.get("content", "")
        history_text += f"{role}: {content}\n"

    prompt = f"""
You are a contract document assistant.

Answer the user's question using only the uploaded documents.
If the answer is not available in the uploaded documents, say that it is not available from the uploaded documents.
Where possible, mention the relevant document name and clause or section context.
Do not provide legal advice. Provide review support only.

Recent chat history:
{history_text}

User question:
{question}

Uploaded documents:
{documents_context}
"""

    return call_model(
        prompt=prompt,
        instructions="You answer questions based only on the uploaded contract documents."
    )


def create_risk_dataframe(analysis):
    rows = analysis.get("risk_register", [])

    if not rows:
        return pd.DataFrame(columns=[
            "Risk Area",
            "Clause Reference",
            "Risk Level",
            "Issue Identified",
            "Business Impact",
            "Recommended Action"
        ])

    data = []

    for item in rows:
        data.append({
            "Risk Area": item.get("risk_area", ""),
            "Clause Reference": item.get("clause_reference", ""),
            "Risk Level": item.get("risk_level", ""),
            "Issue Identified": item.get("issue_identified", ""),
            "Business Impact": item.get("business_impact", ""),
            "Recommended Action": item.get("recommended_action", "")
        })

    return pd.DataFrame(data)


def create_missing_clauses_dataframe(analysis):
    rows = analysis.get("missing_clauses", [])

    if not rows:
        return pd.DataFrame(columns=[
            "Clause Name",
            "Status",
            "Document Reference",
            "Impact",
            "Recommended Action"
        ])

    data = []

    for item in rows:
        data.append({
            "Clause Name": item.get("clause_name", ""),
            "Status": item.get("status", ""),
            "Document Reference": item.get("document_reference", ""),
            "Impact": item.get("impact", ""),
            "Recommended Action": item.get("recommended_action", "")
        })

    return pd.DataFrame(data)


def risk_level_to_score(risk_level):
    value = str(risk_level).strip().lower()

    if value == "high":
        return 3
    if value == "medium":
        return 2
    if value == "low":
        return 1

    return 0


def style_risk_dataframe(df):
    def colour_risk(value):
        value = str(value).strip().lower()

        if value == "high":
            return "background-color: #fecaca; color: #7f1d1d; font-weight: bold;"
        if value == "medium":
            return "background-color: #fde68a; color: #78350f; font-weight: bold;"
        if value == "low":
            return "background-color: #bbf7d0; color: #14532d; font-weight: bold;"

        return ""

    if df.empty:
        return df

    return df.style.applymap(colour_risk, subset=["Risk Level"])


def create_word_file(comparison_text, analysis, file_names):
    document = Document()

    document.add_heading("Forvis Mazars Contracts Analyser", level=0)
    document.add_paragraph(f"Generated on: {datetime.now().strftime('%d %B %Y, %H:%M')}")

    document.add_heading("Uploaded documents", level=1)
    for file_name in file_names:
        document.add_paragraph(file_name, style="List Bullet")

    document.add_heading("Important note", level=1)
    document.add_paragraph(
        "This AI-generated comparison is for review support only. "
        "It should be validated by an appropriate human reviewer before use in any contractual, legal, client-facing or commercial decision."
    )

    document.add_heading("Overall risk summary", level=1)
    summary = analysis.get("summary", {})
    document.add_paragraph(f"Overall risk level: {summary.get('overall_risk_level', 'Not available')}")
    document.add_paragraph(summary.get("overall_comments", "Not available"))

    document.add_heading("Document comparison output", level=1)

    for line in comparison_text.splitlines():
        clean_line = line.strip()

        if not clean_line:
            document.add_paragraph("")
        elif clean_line.startswith("# "):
            document.add_heading(clean_line.replace("# ", "").strip(), level=1)
        elif clean_line.startswith("## "):
            document.add_heading(clean_line.replace("## ", "").strip(), level=2)
        elif clean_line.startswith("### "):
            document.add_heading(clean_line.replace("### ", "").strip(), level=3)
        else:
            document.add_paragraph(clean_line)

    document.add_heading("Missing clause detection", level=1)

    missing_clauses = analysis.get("missing_clauses", [])

    if missing_clauses:
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"

        headers = [
            "Clause Name",
            "Status",
            "Document Reference",
            "Impact",
            "Recommended Action"
        ]

        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header

        for item in missing_clauses:
            row_cells = table.add_row().cells
            row_cells[0].text = str(item.get("clause_name", ""))
            row_cells[1].text = str(item.get("status", ""))
            row_cells[2].text = str(item.get("document_reference", ""))
            row_cells[3].text = str(item.get("impact", ""))
            row_cells[4].text = str(item.get("recommended_action", ""))
    else:
        document.add_paragraph("No missing clause data was generated.")

    document.add_heading("Risk register", level=1)

    risk_register = analysis.get("risk_register", [])

    if risk_register:
        table = document.add_table(rows=1, cols=6)
        table.style = "Table Grid"

        headers = [
            "Risk Area",
            "Clause Reference",
            "Risk Level",
            "Issue Identified",
            "Business Impact",
            "Recommended Action"
        ]

        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header

        for item in risk_register:
            row_cells = table.add_row().cells
            row_cells[0].text = str(item.get("risk_area", ""))
            row_cells[1].text = str(item.get("clause_reference", ""))
            row_cells[2].text = str(item.get("risk_level", ""))
            row_cells[3].text = str(item.get("issue_identified", ""))
            row_cells[4].text = str(item.get("business_impact", ""))
            row_cells[5].text = str(item.get("recommended_action", ""))
    else:
        document.add_paragraph("No risk register data was generated.")

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return buffer.getvalue()


def create_excel_file(comparison_text, analysis, file_names):
    output = io.BytesIO()

    risk_df = create_risk_dataframe(analysis)
    missing_df = create_missing_clauses_dataframe(analysis)

    summary = analysis.get("summary", {})

    summary_df = pd.DataFrame([
        {
            "Generated On": datetime.now().strftime("%d %B %Y, %H:%M"),
            "Uploaded Documents": ", ".join(file_names),
            "Overall Risk Level": summary.get("overall_risk_level", ""),
            "Overall Comments": summary.get("overall_comments", "")
        }
    ])

    comparison_df = pd.DataFrame({
        "Comparison Output": comparison_text.splitlines()
    })

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        risk_df.to_excel(writer, sheet_name="Risk Register", index=False)
        missing_df.to_excel(writer, sheet_name="Missing Clauses", index=False)
        comparison_df.to_excel(writer, sheet_name="Comparison Output", index=False)

    output.seek(0)
    return output.getvalue()


missing_variables = check_environment()

if missing_variables:
    st.error("Missing environment variables: " + ", ".join(missing_variables))
    st.stop()


if "documents" not in st.session_state:
    st.session_state.documents = []

if "comparison_output" not in st.session_state:
    st.session_state.comparison_output = None

if "analysis_output" not in st.session_state:
    st.session_state.analysis_output = None

if "raw_analysis_output" not in st.session_state:
    st.session_state.raw_analysis_output = None

if "uploaded_file_names" not in st.session_state:
    st.session_state.uploaded_file_names = []

if "chat_history" not in st.session_state:
    st.session_state.chat_history = []


with st.sidebar:
    st.header("Configuration")
    #st.write("Azure OpenAI endpoint")
    #st.code(AZURE_OPENAI_ENDPOINT)

    st.write("Model in Use")
    st.code(AZURE_OPENAI_MODEL)

    st.info("Supported files: PDF, DOCX and TXT")

    st.link_button( "🌐 UK Legislation Website", 
                    "https://www.legislation.gov.uk/explore", 
                      use_container_width=True   )
    
    st.link_button( "🌐 Legal Team Useful Documents", 
                    "https://mazarsglobalcloudgbr.sharepoint.com/sites/gbr-hive-sl-quality-risk-management/Shared%20Documents/Forms/AllItems.aspx?FolderCTID=0x01200050BB9F557FCF074EA3ED5CD549C6A0B6&id=%2Fsites%2Fgbr%2Dhive%2Dsl%2Dquality%2Drisk%2Dmanagement%2FShared%20Documents%2FLegal", 
                      use_container_width=True   )


    if st.button("Clear session"):
        st.session_state.documents = [] 
        st.session_state.comparison_output = None
        st.session_state.analysis_output = None
        st.session_state.raw_analysis_output = None
        st.session_state.uploaded_file_names = []
        st.session_state.chat_history = []
        st.session_state.uploader_key += 1
        st.rerun()


upload_tab, compare_tab, risk_tab, chat_tab, download_tab = st.tabs([
    "Upload documents",
    "Compare documents",
    "Missing clauses and risk heatmap",
    "Chat with documents",
    "Download reports"
])


with upload_tab:
    st.markdown("## Upload contract documents")

    st.markdown(
        """
        <div class='info-box'>
        Upload two or more contract documents. Once uploaded, select Extract documents.
        The extracted content will be used for comparison, missing clause detection, risk analysis and document chat.
        </div>
        """,
        unsafe_allow_html=True
    )

    
    uploaded_files = st.file_uploader(
        "Upload contract documents",
        type=["pdf", "docx", "txt"],
        accept_multiple_files=True,
        key=f"contract_uploader_{st.session_state.uploader_key}"
    )


    extract_button = st.button("Extract documents", type="primary", use_container_width=True)

    if extract_button:
        if not uploaded_files:
            st.warning("Please upload at least one document.")
        else:
            with st.spinner("Extracting document text..."):
                extracted_documents = []
                errors = []

                for uploaded_file in uploaded_files:
                    try:
                        extracted_document = extract_text(uploaded_file)
                        extracted_documents.append(extracted_document)
                    except Exception as error:
                        errors.append(f"{uploaded_file.name}: {error}")

                if errors:
                    st.error("Some files could not be processed.")
                    for error in errors:
                        st.write(error)

                st.session_state.documents = extracted_documents
                st.session_state.uploaded_file_names = [
                    document["file_name"] for document in extracted_documents
                ]

                st.success(f"Extracted {len(extracted_documents)} document or documents successfully.")

    if st.session_state.documents:
        st.markdown("## Extracted document preview")

        for document in st.session_state.documents:
            with st.expander(document["file_name"]):
                preview_text = document["text"][:5000]
                st.text_area(
                    label="Extracted text preview",
                    value=preview_text,
                    height=250,
                    key=f"preview_{document['file_name']}"
                )


with compare_tab:
    st.markdown("## Compare documents")

    if len(st.session_state.documents) < 2:
        st.warning("Please upload and extract at least two documents first.")
    else:
        st.write(f"{len(st.session_state.documents)} documents are ready for comparison.")

        compare_button = st.button("Compare documents", type="primary", use_container_width=True)

        if compare_button:
            with st.spinner("Comparing documents using Azure OpenAI..."):
                try:
                    comparison_result = compare_documents(st.session_state.documents)
                    analysis_result, raw_analysis = analyse_missing_clauses_and_risks(st.session_state.documents)

                    st.session_state.comparison_output = comparison_result
                    st.session_state.analysis_output = analysis_result
                    st.session_state.raw_analysis_output = raw_analysis

                    st.success("Comparison, missing clause detection and risk analysis completed.")
                except Exception as error:
                    st.error("The comparison failed.")
                    st.exception(error)

    if st.session_state.comparison_output:
        st.markdown("## Comparison output")
        st.markdown(st.session_state.comparison_output)


with risk_tab:
    st.markdown("## Missing clauses and risk heatmap")

    if not st.session_state.analysis_output:
        st.warning("Run document comparison first to generate missing clause detection and risk heatmap.")
    else:
        analysis = st.session_state.analysis_output
        summary = analysis.get("summary", {})

        col1, col2 = st.columns(2)

        with col1:
            st.metric("Overall risk level", summary.get("overall_risk_level", "Not available"))

        with col2:
            st.write("Overall comments")
            st.info(summary.get("overall_comments", "Not available"))

        st.markdown("## Missing clause detection")
        missing_df = create_missing_clauses_dataframe(analysis)

        if missing_df.empty:
            st.info("No missing clause data was generated.")
        else:
            st.dataframe(missing_df, use_container_width=True)

        st.markdown("## Risk heatmap")
        risk_df = create_risk_dataframe(analysis)

        if risk_df.empty:
            st.info("No risk register data was generated.")
        else:
            high_count = len(risk_df[risk_df["Risk Level"].str.lower() == "high"])
            medium_count = len(risk_df[risk_df["Risk Level"].str.lower() == "medium"])
            low_count = len(risk_df[risk_df["Risk Level"].str.lower() == "low"])

            metric_col1, metric_col2, metric_col3 = st.columns(3)

            with metric_col1:
                st.metric("High risks", high_count)

            with metric_col2:
                st.metric("Medium risks", medium_count)

            with metric_col3:
                st.metric("Low risks", low_count)

            st.dataframe(style_risk_dataframe(risk_df), use_container_width=True)

            heatmap_df = risk_df.copy()
            heatmap_df["Risk Score"] = heatmap_df["Risk Level"].apply(risk_level_to_score)

            chart_df = heatmap_df.groupby("Risk Area", as_index=False)["Risk Score"].max()
            chart_df = chart_df.sort_values("Risk Score", ascending=False)

            st.markdown("## Risk score by area")
            st.bar_chart(
                chart_df,
                x="Risk Area",
                y="Risk Score",
                use_container_width=True
            )


with chat_tab:
    st.markdown("## Chat with uploaded documents")

    if not st.session_state.documents:
        st.warning("Please upload and extract documents first.")
    else:
        st.info("Ask questions based on the uploaded documents. The answer will be grounded only in the extracted document text.")

        for message in st.session_state.chat_history:
            with st.chat_message(message["role"]):
                st.markdown(message["content"])

        user_question = st.chat_input("Ask a question about the uploaded contracts")

        if user_question:
            st.session_state.chat_history.append({
                "role": "user",
                "content": user_question
            })

            with st.chat_message("user"):
                st.markdown(user_question)

            with st.chat_message("assistant"):
                with st.spinner("Reviewing uploaded documents..."):
                    try:
                        answer = chat_with_documents(
                            question=user_question,
                            documents=st.session_state.documents,
                            chat_history=st.session_state.chat_history
                        )

                        st.markdown(answer)

                        st.session_state.chat_history.append({
                            "role": "assistant",
                            "content": answer
                        })

                    except Exception as error:
                        error_message = "The document chat request failed."
                        st.error(error_message)
                        st.exception(error)

                        st.session_state.chat_history.append({
                            "role": "assistant",
                            "content": error_message
                        })


with download_tab:
    st.markdown("## Download reports")

    if not st.session_state.comparison_output or not st.session_state.analysis_output:
        st.warning("Run document comparison first to enable downloads.")
    else:
        word_file = create_word_file(
            comparison_text=st.session_state.comparison_output,
            analysis=st.session_state.analysis_output,
            file_names=st.session_state.uploaded_file_names
        )

        excel_file = create_excel_file(
            comparison_text=st.session_state.comparison_output,
            analysis=st.session_state.analysis_output,
            file_names=st.session_state.uploaded_file_names
        )

        col1, col2 = st.columns(2)

        with col1:
            st.download_button(
                label="Download Word report",
                data=word_file,
                file_name="forvis_mazars_contract_analysis.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )

        with col2:
            st.download_button(
                label="Download Excel risk register",
                data=excel_file,
                file_name="forvis_mazars_contract_risk_register.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

        st.success("Reports are ready for download.")